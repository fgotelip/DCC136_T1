"""Roda construtivo + busca local em todas as instâncias de um zip e preenche a planilha do trabalho.

Exemplo (dentro da pasta src):
    python rodar_lote.py --zip ../drive-download.zip --planilha ../Grupos_para_o_Trabalho_1.xlsx --grupo G4

Não altera nenhum outro arquivo do projeto. Pode ser interrompido e rodado de novo (retoma de onde parou).
"""

import argparse
import csv
import glob
import os
import re
import shutil
import time
import zipfile
from concurrent.futures import ProcessPoolExecutor

import numpy as np
import openpyxl

import construtivo
from busca_local import busca_local
from construtivo import construir
from instancia import ler_instancia
from solucao import calcular_custo, verificar_viabilidade

PADRAO = re.compile(r"instancia_?([12])_?g(\d+)", re.IGNORECASE)


def chave(nome):
    """Instancia1_G7(1).txt -> (1, 7); devolve None se o nome não parece uma instância."""
    achou = PADRAO.search(nome.replace("(1)", ""))
    return (int(achou.group(1)), int(achou.group(2))) if achou else None


def extrair(zip_path, pasta):
    """Extrai o zip (e zips dentro dele) de forma plana para a pasta."""
    os.makedirs(pasta, exist_ok=True)
    fila = [zip_path]
    while fila:
        with zipfile.ZipFile(fila.pop()) as arq:
            for info in arq.infolist():
                if info.is_dir() or "__MACOSX" in info.filename:
                    continue
                destino = os.path.join(pasta, os.path.basename(info.filename))
                with arq.open(info) as origem, open(destino, "wb") as saida:
                    shutil.copyfileobj(origem, saida)
                if destino.lower().endswith(".zip"):
                    fila.append(destino)


def achar_arquivos(pasta):
    """Um arquivo por (tipo, grupo); se houver repetidos fica o maior (ignora arquivos quebrados)."""
    arquivos = {}
    for caminho in glob.glob(os.path.join(pasta, "*.txt")):
        k = chave(os.path.basename(caminho))
        if k and (k not in arquivos or os.path.getsize(caminho) > os.path.getsize(arquivos[k])):
            arquivos[k] = caminho
    return arquivos


def log_progresso(pasta_lote, mensagem):
    caminho_log = os.path.join(os.path.dirname(pasta_lote), "progresso.log")
    horario = time.strftime("%H:%M:%S")
    linha = f"[{horario}] {mensagem}\n"
    print(linha, end="", flush=True)
    try:
        with open(caminho_log, "a", encoding="utf-8") as f:
            f.write(linha)
            f.flush()
    except Exception:
        pass


def rodar_instancia(args):
    nome, caminho, maximos, n_exec, limite_tempo, limite_nos, pasta_lote = args
    construtivo.LIMITE_DE_NOS_DA_BUSCA = limite_nos
    log_progresso(pasta_lote, f"Iniciando {nome} (arquivo: {os.path.basename(caminho)})")
    try:
        ids_c, tipo_c, ids_e, cap_salas, dist = ler_instancia(caminho)
        t0 = time.perf_counter()
        inicial = construir(tipo_c, cap_salas, dist)
        t_cons = time.perf_counter() - t0
        custo_ini = calcular_custo(inicial, dist)
        log_progresso(pasta_lote, f"{nome}: construtivo concluido em {t_cons:.2f}s | custo inicial = {custo_ini:.2f} km")
    except Exception as erro:    # arquivo quebrado ou construtivo sem solução
        log_progresso(pasta_lote, f"{nome}: FALHA no construtivo ({type(erro).__name__}: {erro})")
        for k in maximos:
            with open(os.path.join(pasta_lote, f"{nome}_K{k}.csv"), "w", newline="", encoding="utf-8") as f:
                csv.writer(f).writerows([["status"], [f"FALHA: {type(erro).__name__}: {erro}"]])
        return nome

    for k in maximos:
        linhas = []
        for semente in range(1, n_exec + 1):
            t0 = time.perf_counter()
            sol = busca_local(inicial, tipo_c, cap_salas, dist, k, semente, limite_tempo)
            t_busca = time.perf_counter() - t0
            problemas = verificar_viabilidade(sol, tipo_c, cap_salas, k)
            c_final = calcular_custo(sol, dist)
            linhas.append({"instancia": nome, "k": k, "semente": semente, "custo_inicial": round(custo_ini, 4),
                           "custo_final": round(c_final, 4),
                           "tempo_construtivo": round(t_cons, 4), "tempo_busca": round(t_busca, 4),
                           "tempo_total": round(t_cons + t_busca, 4),
                           "status": "ok" if not problemas else f"INVIAVEL: {problemas[:2]}"})
            log_progresso(pasta_lote, f"{nome} (K={k}) | semente {semente:2d}/{n_exec}: custo = {c_final:10.2f} km | tempo = {t_busca:6.1f}s")
        with open(os.path.join(pasta_lote, f"{nome}_K{k}.csv"), "w", newline="", encoding="utf-8") as f:
            escritor = csv.DictWriter(f, fieldnames=list(linhas[0]))
            escritor.writeheader()
            escritor.writerows(linhas)
        msg_final = (f"{nome} K={k} FINALIZADA: melhor {min(l['custo_final'] for l in linhas):.2f} | "
                     f"media {np.mean([l['custo_final'] for l in linhas]):.2f} | "
                     f"tempo medio {np.mean([l['tempo_total'] for l in linhas]):.2f}s")
        log_progresso(pasta_lote, msg_final)
    return nome


def resumir(pasta_lote):
    """Lê os csv do lote e devolve {(instancia, k): dict com melhor, média, tempo médio, status}."""
    resumo = {}
    for caminho in sorted(glob.glob(os.path.join(pasta_lote, "*.csv"))):
        base = os.path.basename(caminho)[:-4]
        nome, k = base.rsplit("_K", 1)
        with open(caminho, encoding="utf-8") as f:
            linhas = list(csv.DictReader(f))
        if "custo_final" not in linhas[0]:
            resumo[(nome, int(k))] = {"status": linhas[0]["status"]}
            continue
        custos = [float(l["custo_final"]) for l in linhas]
        resumo[(nome, int(k))] = {
            "melhor": min(custos), "media": float(np.mean(custos)),
            "tempo": float(np.mean([float(l["tempo_total"]) for l in linhas])),
            "execucoes": len(linhas),
            "status": "ok" if all(l["status"] == "ok" for l in linhas) else "INVIAVEL"}
    return resumo


def preencher(planilha, saida, grupo, resumo, k_planilha):
    wb = openpyxl.load_workbook(planilha)
    ws = wb["Resultados"]
    coluna = next(c.column for c in ws[1] if c.value == grupo)    # a 1ª coluna do bloco do grupo
    for linha in range(3, ws.max_row + 1):
        nome = ws.cell(linha, 1).value
        r = resumo.get((nome, k_planilha))
        if nome and r and r["status"] == "ok":
            ws.cell(linha, coluna, round(r["melhor"], 2))
            ws.cell(linha, coluna + 1, round(r["media"], 2))
            ws.cell(linha, coluna + 2, round(r["tempo"], 2))
    detalhes = wb.create_sheet(f"{grupo}_detalhes")    # todos os K lado a lado, para conferência
    detalhes.append(["instancia", "K", "melhor", "media", "tempo_medio_s", "execucoes", "status"])
    for (nome, k), r in sorted(resumo.items()):
        detalhes.append([nome, k, r.get("melhor"), r.get("media"), r.get("tempo"), r.get("execucoes"), r["status"]])
    wb.save(saida)


def comparar(planilha, grupo, resumo, k_planilha):
    """Compara os valores já escritos na planilha com os recalculados (melhor e média devem bater)."""
    ws = openpyxl.load_workbook(planilha, data_only=True)["Resultados"]
    coluna = next(c.column for c in ws[1] if c.value == grupo)
    print(f"{'instância':16}{'melhor(plan)':>14}{'melhor(novo)':>14}{'média(plan)':>14}{'média(novo)':>14}  ok?")
    for linha in range(3, ws.max_row + 1):
        nome = ws.cell(linha, 1).value
        r = resumo.get((nome, k_planilha))
        if not nome or not r or r["status"] != "ok":
            continue
        pm, pmed = ws.cell(linha, coluna).value, ws.cell(linha, coluna + 1).value
        bate = pm is not None and abs(pm - r["melhor"]) < 0.01 and abs(pmed - r["media"]) < 0.01
        print(f"{nome:16}{pm!s:>14}{r['melhor']:>14.2f}{pmed!s:>14}{r['media']:>14.2f}  {'sim' if bate else 'DIFERENTE'}")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--zip", required=True)
    ap.add_argument("--planilha", required=True)
    ap.add_argument("--grupo", default="G4")
    ap.add_argument("--max-tipos", type=int, nargs="+", choices=[1, 2], default=[1, 2])
    ap.add_argument("--k-planilha", type=int, choices=[1, 2], default=1, help="K usado nas colunas do grupo")
    ap.add_argument("--execucoes", type=int, default=10)
    ap.add_argument("--limite-de-tempo", type=float, default=300, help="segundos máximos por busca local")
    ap.add_argument("--limite-nos", type=int, default=5_000_000, help="limite do backtracking do construtivo")
    ap.add_argument("--jobs", type=int, default=1, help="instâncias em paralelo (>1 distorce os tempos)")
    ap.add_argument("--saida", default="../resultados_lote")
    ap.add_argument("--so-instancias", nargs="*", help="rodar só estas, ex: Instancia1_G4 Instancia2_G4")
    ap.add_argument("--comparar", help="planilha já preenchida para conferir contra os valores recalculados")
    a = ap.parse_args()

    pasta_inst, pasta_lote = os.path.join(a.saida, "instancias"), os.path.join(a.saida, "lote")
    os.makedirs(pasta_lote, exist_ok=True)
    if not os.path.isdir(pasta_inst):
        extrair(a.zip, pasta_inst)
    arquivos = achar_arquivos(pasta_inst)

    tarefas = []
    for (tipo, g), caminho in sorted(arquivos.items(), key=lambda x: (x[0][0], x[0][1])):
        nome = f"Instancia{tipo}_G{g}"
        if a.so_instancias and nome not in a.so_instancias:
            continue
        faltam = [k for k in a.max_tipos if not os.path.exists(os.path.join(pasta_lote, f"{nome}_K{k}.csv"))]
        if faltam:    # retoma: pula o que já foi calculado
            tarefas.append((nome, caminho, faltam, a.execucoes, a.limite_de_tempo, a.limite_nos, pasta_lote))
        print(f"{nome:16} <- {os.path.basename(caminho)}")

    if a.jobs > 1:
        with ProcessPoolExecutor(a.jobs) as pool:
            list(pool.map(rodar_instancia, tarefas))
    else:
        for t in tarefas:
            rodar_instancia(t)

    resumo = resumir(pasta_lote)
    if a.comparar:
        comparar(a.comparar, a.grupo, resumo, a.k_planilha)
    else:
        saida = os.path.join(a.saida, f"planilha_{a.grupo}_preenchida.xlsx")
        preencher(a.planilha, saida, a.grupo, resumo, a.k_planilha)
        print("planilha salva em", saida)
    falhas = [(n, k, r["status"]) for (n, k), r in resumo.items() if r["status"] != "ok"]
    for n, k, s in falhas:
        print(f"SEM RESULTADO {n} K={k}: {s}")
